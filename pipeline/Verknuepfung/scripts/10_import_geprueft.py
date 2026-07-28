# -*- coding: utf-8 -*-
"""
10_import_geprueft.py
Übernimmt manuell GEPRÜFTE Lärmquellen (Spalte Laermquelle_geprueft) aus einer
Backup-Arbeitsmappe in relevante_ereignisse.csv. Mapping per WAV-Dateiname.
Quelle (Default): Laermquellen_Verknuepfung_backup.xlsx, alle Blätter mit den
Spalten WAV + Laermquelle_geprueft (insb. 'Pruefen_NichtBaulaerm' und 'Ereignisse').
"""
import os, csv, sys
import openpyxl

HERE = os.path.dirname(__file__)
VK   = os.path.abspath(os.path.join(HERE, ".."))
ALIAS = {"Tiefbohrer": "Bohrgeraet/schweres Geraet"}   # gleiche Maschine -> zusammenführen
SRC  = sys.argv[1] if len(sys.argv) > 1 else os.path.join(VK, "Laermquellen_Verknuepfung_backup.xlsx")
TGT  = os.path.join(VK, "relevante_ereignisse.csv")

# ---- geprüfte Labels aus Backup sammeln (WAV -> geprueft)
checked = {}
wb = openpyxl.load_workbook(SRC, read_only=True)
for sn in wb.sheetnames:
    ws = wb[sn]
    hdr = [c.value for c in next(ws.iter_rows(max_row=1))]
    if "WAV" in hdr and "Laermquelle_geprueft" in hdr:
        wi, gi = hdr.index("WAV"), hdr.index("Laermquelle_geprueft")
        for row in ws.iter_rows(min_row=2, values_only=True):
            wav = row[wi] if wi < len(row) else None
            g   = row[gi] if gi < len(row) else None
            if wav and g not in (None, ""):
                lab = str(g).strip()
                checked[str(wav)] = ALIAS.get(lab, lab)   # "Tiefbohrer" -> Bohrgerät
print(f"Geprüfte Labels im Backup: {len(checked)}")

# ---- in relevante_ereignisse.csv eintragen (vorhandene geprueft NICHT überschreiben)
rows = list(csv.DictReader(open(TGT, newline='', encoding='utf-8-sig')))
fields = list(rows[0].keys())
if "Laermquelle_geprueft" not in fields:
    fields.insert(fields.index("Laermquelle_KI")+1 if "Laermquelle_KI" in fields else len(fields),
                  "Laermquelle_geprueft")
upd = 0
for r in rows:
    cur = (r.get("Laermquelle_geprueft") or "").strip()
    new = checked.get(r["WAV"])
    if new and not cur:
        r["Laermquelle_geprueft"] = new
        upd += 1
with open(TGT, "w", newline='', encoding='utf-8-sig') as f:
    w = csv.DictWriter(f, fieldnames=fields); w.writeheader()
    for r in rows: w.writerow({k: r.get(k, "") for k in fields})

from collections import Counter
print(f"In relevante_ereignisse.csv übernommen: {upd}")
print("Verteilung geprueft:", dict(Counter(r.get("Laermquelle_geprueft","") or "(leer)" for r in rows)))
